"""
build_office.py — DOCX and XLSX members of the synthetic MRPL corpus.

The spreadsheets are not decoration. They are the numeric substrate the agent
runs real calculations over in the sandbox demo (vibration trend slope, spares
coverage, downtime cost), so the numbers are internally consistent with the
narrative documents.
"""

import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import corpus_spec as S  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor, Cm  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import (Alignment, Border, Font, PatternFill,  # noqa: E402
                             Side)
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.chart import LineChart, Reference  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "documents"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "0B3D5C"
HDR_FILL = PatternFill("solid", fgColor="DCE6EF")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
WARN_FILL = PatternFill("solid", fgColor="FFE699")
OK_FILL = PatternFill("solid", fgColor="D9EAD3")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
# DOCX: Management of Change
# ---------------------------------------------------------------------------
def _set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _kv_table(doc, rows):
    t = doc.add_table(rows=0, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.bold = (i % 2 == 0)
            if i % 2 == 0:
                _set_cell_bg(cells[i], "EEF2F6")
    return t


def _grid(doc, header, rows, widths=None, fontsize=8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = h
        _set_cell_bg(c, "DCE6EF")
        for p in c.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(fontsize)
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(fontsize)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def build_moc():
    d = S.DOCS["moc"]
    M = S.MOC
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"{S.PLANT['short']}  |  {S.unit_header()}  |  {d['id']}")
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string(NAVY)
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"{S.PLANT['doc_class']}   -   {S.PLANT['synthetic_notice']}")
    fr.font.size = Pt(7)
    fr.font.color.rgb = RGBColor.from_string("777777")
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("MANAGEMENT OF CHANGE REQUEST")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sp.add_run(f"{d['id']}  |  {d['title']}")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor.from_string("444444")

    _kv_table(doc, [
        ("MOC Number", d["id"], "Date raised", d["raised"]),
        ("Change type", M["change_type"], "Risk rank", M["risk_rank"]),
        ("Unit", f"{S.PLANT['unit_code']} - {S.PLANT['unit_name']}",
         "Equipment", "P-2104A / P-2104B"),
        ("Originator", S.PEOPLE["rot_eng"], "Current status", d["status"]),
        ("Target completion", d["target_completion"], "Arising from",
         S.DOCS["incident"]["id"]),
    ])

    doc.add_heading("1. Reason for change", level=2)
    doc.add_paragraph(M["driver"])

    doc.add_heading("2. Existing arrangement", level=2)
    doc.add_paragraph(M["existing"])

    doc.add_heading("3. Proposed arrangement", level=2)
    doc.add_paragraph(M["proposed"])

    doc.add_heading("4. Technical justification", level=2)
    for j in M["justification"]:
        doc.add_paragraph(j, style="List Bullet")

    doc.add_heading("5. Impact assessment", level=2)
    _grid(doc, ["Discipline", "Impact"], M["impacts"], widths=[3.2, 13.3])

    doc.add_heading("6. Constraint on execution", level=2)
    p = doc.add_paragraph()
    r = p.add_run(M["blocker"])
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("8A1C1C")
    doc.add_paragraph(
        "Consequently the earliest realistic execution date is governed by "
        "material availability rather than by engineering approval. The 45 day "
        f"vendor lead time on seal kit SK-682-2104-A3 means that even with "
        f"immediate approval and a purchase order placed on {d['raised']}, "
        "material would not be on site before mid October 2026. The target "
        f"completion date of {d['target_completion']} should be regarded as at "
        "risk."
    )

    doc.add_heading("7. Safety review comment", level=2)
    p = doc.add_paragraph()
    p.add_run(f"{S.PEOPLE['safety_officer']} (Safety Officer), 2026-09-03: ").bold = True
    p.add_run(M["safety_comment"])

    doc.add_heading("8. Pre-commissioning requirements", level=2)
    _grid(doc, ["No.", "Requirement", "Responsible"], [
        ("1", "Revise SOP-CDU2-014 to Rev 5 covering barrier fluid pre-start "
              "checks and the 0.05 mm alignment tolerance record.",
         S.PEOPLE["dgm_ops"]),
        ("2", "Configure two new DCS alarm points for barrier fluid level low "
              "and barrier fluid pressure low.", "Instrumentation"),
        ("3", "Complete 2 hour refresher for all CDU-2 shift personnel.",
         S.PEOPLE["dgm_ops"]),
        ("4", "Update the P&ID PID-CDU2-003 and the equipment data sheet to "
              "reflect the Plan 53B system.", "Engineering"),
        ("5", "Add barrier fluid level to the operator shift round check "
              "sheet (also required by CAPA-0042-3).", S.PEOPLE["panel_op"]),
        ("6", "Obtain hot work deviation permit or schedule tie-in inside the "
              "planned shutdown window, per SC-2026-11.", S.PEOPLE["planner"]),
    ], widths=[1.2, 11.5, 3.8])

    doc.add_heading("9. Approval route", level=2)
    _grid(doc, ["Stage", "Name", "Date", "Status"], M["approvals"],
          widths=[5.5, 4.0, 3.0, 4.0])

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(
        "This change shall not be executed until all approval stages above "
        "show Signed and the material availability constraint in Section 6 is "
        "resolved."
    )
    r.font.bold = True
    r.font.size = Pt(9)

    path = OUT / f"{d['id']}_Seal_Upgrade_P-2104.docx"
    doc.core_properties.title = d["title"]
    doc.core_properties.author = S.PEOPLE["rot_eng"]
    doc.core_properties.comments = S.PLANT["synthetic_notice"]
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# XLSX helpers
# ---------------------------------------------------------------------------
def _style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, size=9, color="0B3D5C")
        cell.fill = HDR_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autofit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title_block(ws, doc_id, title, extra=None, ncols=6):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=f"{S.PLANT['short']}  -  {title}")
    c.font = Font(bold=True, size=12, color="0B3D5C")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    sub = f"{S.unit_header()}   |   Document: {doc_id}"
    if extra:
        sub += f"   |   {extra}"
    ws.cell(row=2, column=1, value=sub).font = Font(size=8, color="555555")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    ws.cell(row=3, column=1,
            value=f"{S.PLANT['doc_class']}  -  {S.PLANT['synthetic_notice']}"
            ).font = Font(size=7, italic=True, color="8A1C1C")
    return 5  # first header row


# ---------------------------------------------------------------------------
# XLSX 1: vibration log  (+ embedded trend chart)
# ---------------------------------------------------------------------------
def build_vib_xlsx():
    d = S.DOCS["vib"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Vibration Log"

    hr = _title_block(ws, d["id"], d["title"],
                      extra=f"Limits per {S.VIB_LIMITS['standard']}", ncols=8)
    header = ["Month", "Equipment Tag", "Description",
              "Overall Vibration (mm/s RMS)", "Alert Limit (mm/s)",
              "Trip Limit (mm/s)", "Bearing Temp (degC)", "Status / Remark"]
    for i, h in enumerate(header, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, len(header), row=hr)

    desc = {e["tag"]: e["desc"] for e in S.EQUIPMENT}
    r = hr + 1
    order = ["P-2104B", "P-2104A", "P-2101A", "P-2101B", "P-2107A"]
    for tag in order:
        for month, vib, btemp in S.VIB_READINGS[tag]:
            if vib >= S.VIB_LIMITS["trip_mm_s"]:
                status = "TRIP LEVEL - remove from service"
            elif vib >= S.VIB_LIMITS["alert_mm_s"]:
                status = "ALERT - corrective work order required within 24 h"
            elif vib > S.VIB_LIMITS["normal_max_mm_s"]:
                status = "Above normal band - increase monitoring frequency"
            else:
                status = "Normal"
            if tag == "P-2104B" and month == "2026-07":
                status += " | Signed off as 'monitor'. No work order raised."
            if tag == "P-2104B" and month == "2026-08":
                status += (" | Seal failure 2026-08-29, refer "
                           f"{S.DOCS['incident']['id']}")
            vals = [month, tag, desc.get(tag, ""), vib,
                    S.VIB_LIMITS["alert_mm_s"], S.VIB_LIMITS["trip_mm_s"],
                    btemp, status]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.border = BOX
                cell.font = Font(size=9)
                if i in (4, 5, 6, 7):
                    cell.alignment = Alignment(horizontal="center")
                if i == 8:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    last = r - 1

    ws.conditional_formatting.add(
        f"D{hr+1}:D{last}",
        CellIsRule(operator="greaterThanOrEqual",
                   formula=[str(S.VIB_LIMITS["alert_mm_s"])], fill=BAD_FILL))
    ws.conditional_formatting.add(
        f"D{hr+1}:D{last}",
        CellIsRule(operator="between",
                   formula=[str(S.VIB_LIMITS["normal_max_mm_s"]),
                            str(S.VIB_LIMITS["alert_mm_s"])], fill=WARN_FILL))
    ws.conditional_formatting.add(
        f"G{hr+1}:G{last}",
        CellIsRule(operator="greaterThanOrEqual",
                   formula=[str(S.VIB_LIMITS["bearing_temp_alert_c"])],
                   fill=BAD_FILL))
    _autofit(ws, [10, 14, 42, 15, 12, 12, 13, 58])

    # Trend sheet: pivot-style, one column per tag, with a chart
    ws2 = wb.create_sheet("Trend P-2104")
    hr2 = _title_block(ws2, d["id"], "P-2104A/B Vibration Trend 2026", ncols=5)
    for i, h in enumerate(["Month", "P-2104B (mm/s)", "P-2104A (mm/s)",
                           "Alert Limit", "Trip Limit"], 1):
        ws2.cell(row=hr2, column=i, value=h)
    _style_header(ws2, 5, row=hr2)
    months = [m for m, _, _ in S.VIB_READINGS["P-2104B"]]
    for j, m in enumerate(months):
        row = hr2 + 1 + j
        ws2.cell(row=row, column=1, value=m).border = BOX
        ws2.cell(row=row, column=2,
                 value=S.VIB_READINGS["P-2104B"][j][1]).border = BOX
        ws2.cell(row=row, column=3,
                 value=S.VIB_READINGS["P-2104A"][j][1]).border = BOX
        ws2.cell(row=row, column=4,
                 value=S.VIB_LIMITS["alert_mm_s"]).border = BOX
        ws2.cell(row=row, column=5,
                 value=S.VIB_LIMITS["trip_mm_s"]).border = BOX
    _autofit(ws2, [12, 16, 16, 12, 12])

    ch = LineChart()
    ch.title = "CDU-2 Column Bottoms Pump Vibration Trend 2026"
    ch.y_axis.title = "Overall velocity mm/s RMS"
    ch.x_axis.title = "Month"
    ch.height, ch.width = 8, 18
    data = Reference(ws2, min_col=2, max_col=5, min_row=hr2,
                     max_row=hr2 + len(months))
    cats = Reference(ws2, min_col=1, min_row=hr2 + 1, max_row=hr2 + len(months))
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ws2.add_chart(ch, f"G{hr2}")

    # Limits reference sheet
    ws3 = wb.create_sheet("Limits")
    hr3 = _title_block(ws3, d["id"], "Condition Monitoring Limits", ncols=3)
    for i, h in enumerate(["Parameter", "Value", "Basis"], 1):
        ws3.cell(row=hr3, column=i, value=h)
    _style_header(ws3, 3, row=hr3)
    lim_rows = [
        ("Normal maximum vibration (mm/s RMS)",
         S.VIB_LIMITS["normal_max_mm_s"], S.VIB_LIMITS["standard"]),
        ("Alert vibration (mm/s RMS)", S.VIB_LIMITS["alert_mm_s"],
         f"{S.VIB_LIMITS['standard']} | {S.DOCS['std']['id']} clause 5.1"),
        ("Trip vibration (mm/s RMS)", S.VIB_LIMITS["trip_mm_s"],
         f"{S.VIB_LIMITS['standard']} | {S.DOCS['std']['id']} clause 5.1"),
        ("Bearing temperature alarm (degC)",
         S.VIB_LIMITS["bearing_temp_alert_c"],
         f"{S.DOCS['std']['id']} clause 5.2"),
        ("Bearing temperature trip (degC)",
         S.VIB_LIMITS["bearing_temp_trip_c"],
         f"{S.DOCS['std']['id']} clause 5.2"),
        ("Escalation rule", "Work order within 24 hours of an alert reading",
         f"{S.DOCS['std']['id']} clause 5.1"),
    ]
    for j, row in enumerate(lim_rows):
        for i, v in enumerate(row, 1):
            c = ws3.cell(row=hr3 + 1 + j, column=i, value=v)
            c.border = BOX
            c.font = Font(size=9)
    _autofit(ws3, [40, 44, 60])

    path = OUT / f"{d['id']}_Vibration_Log.xlsx"
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# XLSX 2: work order history
# ---------------------------------------------------------------------------
def build_wo_xlsx():
    d = S.DOCS["wo"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Work Orders"
    hr = _title_block(ws, d["id"], d["title"], ncols=9)
    header = ["Work Order", "Equipment Tag", "Date", "Type", "Description",
              "Downtime (hr)", "Labour (hr)", "Cost (INR)", "Status"]
    for i, h in enumerate(header, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, len(header), row=hr)
    for j, row in enumerate(S.WORK_ORDERS):
        r = hr + 1 + j
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=9)
            if i == 5:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 8:
                c.number_format = "#,##0"
            if i in (6, 7):
                c.alignment = Alignment(horizontal="center")
        if row[3] == "Breakdown":
            for i in range(1, len(header) + 1):
                ws.cell(row=r, column=i).fill = BAD_FILL
    last = hr + len(S.WORK_ORDERS)
    r = last + 2
    ws.cell(row=r, column=5, value="TOTAL").font = Font(bold=True, size=9)
    ws.cell(row=r, column=6, value=f"=SUM(F{hr+1}:F{last})").font = Font(bold=True, size=9)
    ws.cell(row=r, column=7, value=f"=SUM(G{hr+1}:G{last})").font = Font(bold=True, size=9)
    c = ws.cell(row=r, column=8, value=f"=SUM(H{hr+1}:H{last})")
    c.font = Font(bold=True, size=9)
    c.number_format = "#,##0"
    ws.cell(row=r + 1, column=5,
            value="TOTAL for P-2104B only").font = Font(bold=True, size=9)
    c = ws.cell(row=r + 1, column=8,
                value=f'=SUMIF(B{hr+1}:B{last},"P-2104B",H{hr+1}:H{last})')
    c.font = Font(bold=True, size=9)
    c.number_format = "#,##0"
    _autofit(ws, [15, 14, 12, 13, 62, 13, 12, 14, 12])
    path = OUT / f"{d['id']}_Work_Order_History.xlsx"
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# XLSX 3: spares inventory  (the procurement blocker lives here)
# ---------------------------------------------------------------------------
def build_spares_xlsx():
    d = S.DOCS["spares"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Spares"
    hr = _title_block(ws, d["id"], d["title"],
                      extra="Stock position as at 2026-09-03", ncols=12)
    header = ["Part Number", "Description", "UoM", "On Hand", "Min Level",
              "Max Level", "Below Min?", "Shortfall", "Lead Time (days)",
              "Vendor", "Unit Cost (INR)", "Criticality", "Applies To"]
    for i, h in enumerate(header, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, len(header), row=hr)
    for j, (part, desc, uom, on_hand, mn, mx, lead, vendor, cost, crit,
            applies) in enumerate(S.SPARES):
        r = hr + 1 + j
        vals = [part, desc, uom, on_hand, mn, mx,
                f"=IF(D{r}<E{r},\"YES\",\"NO\")",
                f"=MAX(0,E{r}-D{r})", lead, vendor, cost, crit, applies]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=9)
            if i == 2:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 11:
                c.number_format = "#,##0"
            if i in (3, 4, 5, 6, 7, 8, 9, 12):
                c.alignment = Alignment(horizontal="center")
        if on_hand < mn:
            for i in range(1, len(header) + 1):
                ws.cell(row=r, column=i).fill = BAD_FILL
        elif on_hand == mn:
            for i in range(1, len(header) + 1):
                ws.cell(row=r, column=i).fill = WARN_FILL
        else:
            ws.cell(row=r, column=4).fill = OK_FILL
    last = hr + len(S.SPARES)

    ws2 = wb.create_sheet("Procurement Risk")
    hr2 = _title_block(ws2, d["id"], "Criticality A Items Below Minimum",
                       ncols=6)
    for i, h in enumerate(["Part Number", "Description", "Shortfall",
                           "Lead Time (days)", "Earliest Availability",
                           "Impact"], 1):
        ws2.cell(row=hr2, column=i, value=h)
    _style_header(ws2, 6, row=hr2)
    risk = [
        ("SK-682-2104-A3",
         "Cartridge mech seal kit, API 682 Cat 2 Arr 3 dual pressurised",
         2, 45, "2026-10-16",
         f"BLOCKS {S.DOCS['moc']['id']}. Target completion "
         f"{S.DOCS['moc']['target_completion']} is at risk. This is the "
         "primary corrective action from " + S.DOCS["incident"]["id"] + "."),
        ("ACC-53B-09", "Bladder accumulator, Plan 53B, 8 litre, 40 bar",
         1, 60, "2026-10-31",
         f"BLOCKS {S.DOCS['moc']['id']}. Longest lead item in the change."),
        ("SK-682-2104",
         "Cartridge mech seal kit, API 682 Cat 1 Arr 1 (existing type)",
         1, 45, "2026-10-16",
         "One kit consumed by WO-26-15118 on 2026-08-29. Site now holds 1 "
         f"against a minimum of 2, contrary to {S.DOCS['std']['id']} "
         "clause 6.1. No spare cover if the second pump fails."),
    ]
    for j, row in enumerate(risk):
        for i, v in enumerate(row, 1):
            c = ws2.cell(row=hr2 + 1 + j, column=i, value=v)
            c.border = BOX
            c.font = Font(size=9)
            c.fill = BAD_FILL
            if i in (2, 6):
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if i in (3, 4, 5):
                c.alignment = Alignment(horizontal="center")
    ws2.cell(row=hr2 + len(risk) + 2, column=1,
             value="Note: earliest availability assumes a purchase order "
                   "placed on 2026-09-01 and no expediting."
             ).font = Font(size=8, italic=True)
    _autofit(ws, [18, 46, 7, 9, 10, 10, 10, 10, 12, 24, 14, 11, 16])
    _autofit(ws2, [18, 44, 10, 14, 16, 62])
    path = OUT / f"{d['id']}_Spares_Inventory.xlsx"
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# XLSX 4: training / competency matrix
# ---------------------------------------------------------------------------
def build_training_xlsx():
    d = S.DOCS["training"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Competency Matrix"
    hr = _title_block(ws, d["id"], d["title"],
                      extra=f"Validity assessed as at {S.TRAINING_AS_OF}",
                      ncols=8)
    header = ["Name", "Role", "Pump Changeover - Hot Standby (expiry)",
              "Hot Work Permit (expiry)", "Confined Space Entry (expiry)",
              "DCS Operation (expiry)", "Last Refresher",
              "Overall Status"]
    for i, h in enumerate(header, 1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, len(header), row=hr)

    from datetime import date
    asof = date.fromisoformat(S.TRAINING_AS_OF)

    for j, row in enumerate(S.TRAINING):
        r = hr + 1 + j
        expired = []
        for idx, col in enumerate(row[2:6]):
            if col == "-":
                continue
            if date.fromisoformat(col) < asof:
                expired.append(header[2 + idx].split(" (")[0])
        status = "VALID" if not expired else "LAPSED: " + "; ".join(expired)
        vals = list(row) + [status]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=9)
            if i in range(3, 8):
                c.alignment = Alignment(horizontal="center")
        for idx in range(4):
            col = row[2 + idx]
            cell = ws.cell(row=r, column=3 + idx)
            if col != "-" and date.fromisoformat(col) < asof:
                cell.fill = BAD_FILL
                cell.font = Font(size=9, bold=True, color="8A1C1C")
        sc = ws.cell(row=r, column=8)
        sc.fill = BAD_FILL if expired else OK_FILL
        sc.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0] == S.PEOPLE["shift_incharge"]:
            ws.cell(row=r, column=1).font = Font(size=9, bold=True)

    r = hr + len(S.TRAINING) + 2
    ws.cell(row=r, column=1, value="NOTES").font = Font(bold=True, size=9)
    notes = [
        f"{S.PEOPLE['shift_incharge']} was the Shift In-Charge on duty during "
        f"incident {S.DOCS['incident']['id']} on {S.DOCS['incident']['date']} "
        "with a Pump Changeover certification that expired 2026-07-31. This is "
        "recorded as a contributing factor in the investigation report.",
        "2 of 6 personnel holding a Shift In-Charge role currently show a "
        "lapsed Pump Changeover - Hot Standby certification.",
        "CAPA-0042-5 requires re-certification of all CDU-2 shift in-charges "
        "and a roster block on lapsed certification, due 2026-09-30.",
        "The roster system does not currently validate certification expiry "
        "before assigning a shift in-charge.",
    ]
    for k, n in enumerate(notes):
        ws.merge_cells(start_row=r + 1 + k, start_column=1,
                       end_row=r + 1 + k, end_column=8)
        c = ws.cell(row=r + 1 + k, column=1, value=f"{k+1}. {n}")
        c.font = Font(size=8)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r + 1 + k].height = 26
    _autofit(ws, [18, 26, 18, 16, 18, 15, 14, 34])
    path = OUT / f"{d['id']}_Competency_Matrix.xlsx"
    wb.save(str(path))
    return path
